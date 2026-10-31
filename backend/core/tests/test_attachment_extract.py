"""Tests for bounded attachment text extraction (backend/core/attachment_extract.py).

Every fixture is generated here; nothing binary is committed.
"""

import inspect
import io
import struct
import sys
import time
import tracemalloc
import zipfile
import zlib

import docx
import pytest
from pypdf import PdfWriter

from backend.core import attachment_extract as ae
from backend.core.attachment_extract import ExtractionLimits, extract

# --- Fixture builders --------------------------------------------------------


def _pdf(pages: list[str]) -> bytes:
    """A PDF with one Helvetica text line per page ("" gives a page with no text)."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog = add(b"")
    pages_id = add(b"")
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    kids = []
    for text in pages:
        stream = (
            f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1") if text else b""
        )
        contents = add(
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        )
        kids.append(
            add(
                b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
                % (pages_id, font, contents)
            )
        )
    objects[catalog - 1] = b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id
    objects[pages_id - 1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(b"%d 0 R" % k for k in kids),
        len(kids),
    )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, catalog, xref)
    )
    return out.getvalue()


def _blank_pdf(pages: int = 2) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(612, 792)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def _zip(members: dict, compression=zipfile.ZIP_DEFLATED) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression) as archive:
        for name, body in members.items():
            archive.writestr(name, body)
    return out.getvalue()


def _docx(paragraphs=("Hello world",), table=None) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


_S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
_REL_WS = _R + "/worksheet"
_REL_SST = _R + "/sharedStrings"
_REL_SLIDE = _R + "/slide"


def _rels(entries) -> str:
    body = "".join(
        f'<Relationship Id="{i}" Type="{t}" Target="{target}"/>'
        for i, t, target in entries
    )
    return f'<?xml version="1.0"?><Relationships xmlns="{_PKG}">{body}</Relationships>'


def _xlsx(
    sheets: dict, shared=("shared one",), doctype: str = "", entity_ref: str = ""
) -> bytes:
    """``sheets`` maps sheet name -> rows of cells; a cell "s:N" is a shared-string ref."""
    workbook_sheets = "".join(
        f'<sheet name="{name}" sheetId="{i}" r:id="rId{i}"/>'
        for i, name in enumerate(sheets, 1)
    )
    members = {
        "[Content_Types].xml": "<Types/>",
        "xl/workbook.xml": (
            f'<?xml version="1.0"?><workbook xmlns="{_S}" xmlns:r="{_R}">'
            f"<sheets>{workbook_sheets}</sheets></workbook>"
        ),
        "xl/_rels/workbook.xml.rels": _rels(
            [
                (f"rId{i}", _REL_WS, f"worksheets/sheet{i}.xml")
                for i in range(1, len(sheets) + 1)
            ]
            + [("rIdS", _REL_SST, "sharedStrings.xml")]
        ),
        "xl/sharedStrings.xml": (
            f'<?xml version="1.0"?>{doctype}<sst xmlns="{_S}">'
            + "".join(f"<si><t>{s}</t></si>" for s in shared)
            + (f"<si><t>{entity_ref}</t></si>" if entity_ref else "")
            + "</sst>"
        ),
    }
    for i, rows in enumerate(sheets.values(), 1):
        xml_rows = []
        for row in rows:
            cells = []
            for value in row:
                if value.startswith("s:"):
                    cells.append(f'<c t="s"><v>{value[2:]}</v></c>')
                elif value.startswith("i:"):
                    cells.append(f'<c t="inlineStr"><is><t>{value[2:]}</t></is></c>')
                else:
                    cells.append(f"<c><v>{value}</v></c>")
            xml_rows.append(f"<row>{''.join(cells)}</row>")
        members[f"xl/worksheets/sheet{i}.xml"] = (
            f'<?xml version="1.0"?><worksheet xmlns="{_S}"><sheetData>'
            f"{''.join(xml_rows)}</sheetData></worksheet>"
        )
    return _zip(members)


_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _pptx(slides: list[list[str]], doctype: str = "") -> bytes:
    """Slides are stored under reversed file numbers so presentation order != file order."""
    count = len(slides)
    members = {
        "[Content_Types].xml": "<Types/>",
        "ppt/presentation.xml": (
            f'<?xml version="1.0"?><p:presentation xmlns:p="{_P}" xmlns:r="{_R}"><p:sldIdLst>'
            + "".join(
                f'<p:sldId id="{255 + i}" r:id="rId{i}"/>' for i in range(1, count + 1)
            )
            + "</p:sldIdLst></p:presentation>"
        ),
        "ppt/_rels/presentation.xml.rels": _rels(
            [
                (f"rId{i}", _REL_SLIDE, f"slides/slide{count + 1 - i}.xml")
                for i in range(1, count + 1)
            ]
        ),
    }
    for i, paragraphs in enumerate(slides, 1):
        body = "".join(f"<a:p><a:r><a:t>{p}</a:t></a:r></a:p>" for p in paragraphs)
        members[f"ppt/slides/slide{count + 1 - i}.xml"] = (
            f'<?xml version="1.0"?>{doctype}<p:sld xmlns:p="{_P}" xmlns:a="{_A}"><p:cSld><p:spTree>'
            f"<p:sp><p:txBody>{body}</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
        )
    return _zip(members)


_ODF_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"'
)


def _odf(mimetype: str, body: str, doctype: str = "") -> bytes:
    content = (
        f'<?xml version="1.0"?>{doctype}<office:document-content {_ODF_NS}>'
        f"<office:body>{body}</office:body></office:document-content>"
    )
    return _zip({"mimetype": mimetype, "content.xml": content})


def _texts(result) -> list[str]:
    return [segment.text for segment in result.segments]


def _total_chars(result) -> int:
    return sum(len(segment.text) for segment in result.segments)


# --- Per-format extraction ---------------------------------------------------


def test_pdf_text_is_returned_per_page_with_page_locators():
    result = extract(_pdf(["First page", "Second page"]), ae.PDF)
    assert result.status == ae.STATUS_OK
    assert [s.locator for s in result.segments] == [{"page": 1}, {"page": 2}]
    assert "First page" in result.segments[0].text
    assert "Second page" in result.segments[1].text


def test_docx_paragraphs_and_table_cells_are_extracted_in_body_order():
    data = _docx(paragraphs=["Intro text"], table=[["Name", "Role"], ["Ada", "Lead"]])
    result = extract(data, ae.DOCX)
    assert result.status == ae.STATUS_OK
    assert result.segments[0].text == "Intro text"
    assert result.segments[0].locator == {"paragraph": 1}
    assert result.segments[1].locator == {"table": 1}
    assert result.segments[1].text == "Name\tRole\nAda\tLead"


def test_xlsx_resolves_shared_inline_and_literal_cells_per_named_sheet():
    data = _xlsx(
        {"Budget": [["s:0", "42"], ["i:inline value"]], "Notes": [["s:1"]]},
        shared=("shared one", "shared two"),
    )
    result = extract(data, ae.XLSX)
    assert result.status == ae.STATUS_OK
    assert [s.locator for s in result.segments] == [
        {"sheet": "Budget"},
        {"sheet": "Notes"},
    ]
    assert result.segments[0].text == "shared one\t42\ninline value"
    assert result.segments[1].text == "shared two"


def test_pptx_slides_follow_presentation_order_not_file_names():
    result = extract(_pptx([["Title slide"], ["Second", "bullet"]]), ae.PPTX)
    assert result.status == ae.STATUS_OK
    assert [s.locator for s in result.segments] == [{"slide": 1}, {"slide": 2}]
    assert _texts(result) == ["Title slide", "Second\nbullet"]


def test_odt_paragraphs_expand_space_and_tab_elements():
    body = (
        "<office:text><text:h>Heading</text:h>"
        '<text:p>two<text:s text:c="2"/>spaces<text:tab/>tab</text:p></office:text>'
    )
    result = extract(_odf(ae.ODT, body), ae.ODT)
    assert result.status == ae.STATUS_OK
    assert _texts(result) == ["Heading", "two  spaces\ttab"]
    assert [s.locator for s in result.segments] == [{"paragraph": 1}, {"paragraph": 2}]


def test_ods_rows_are_returned_per_named_table():
    body = (
        '<office:spreadsheet><table:table table:name="Data">'
        "<table:table-row><table:table-cell><text:p>a</text:p></table:table-cell>"
        "<table:table-cell><text:p>b</text:p></table:table-cell></table:table-row>"
        "</table:table></office:spreadsheet>"
    )
    result = extract(_odf(ae.ODS, body), ae.ODS)
    assert result.status == ae.STATUS_OK
    assert result.segments == (ae.Segment("a\tb", {"sheet": "Data"}),)


def test_odp_text_is_returned_per_slide():
    body = (
        "<office:presentation><draw:page><text:p>One</text:p></draw:page>"
        "<draw:page><text:p>Two</text:p></draw:page></office:presentation>"
    )
    result = extract(_odf(ae.ODP, body), ae.ODP)
    assert [s.locator for s in result.segments] == [{"slide": 1}, {"slide": 2}]
    assert _texts(result) == ["One", "Two"]


def test_text_blocks_carry_their_starting_line_and_bom_is_dropped():
    data = "﻿alpha\nbeta\n\n\ngamma\n".encode()
    result = extract(data, ae.TEXT)
    assert result.status == ae.STATUS_OK
    assert result.segments == (
        ae.Segment("alpha\nbeta", {"line": 1}),
        ae.Segment("gamma", {"line": 5}),
    )


def test_long_csv_is_split_into_bounded_line_blocks():
    data = "".join(f"row{i},value\n" for i in range(1, 121)).encode()
    result = extract(data, ae.CSV)
    assert [s.locator for s in result.segments] == [
        {"line": 1},
        {"line": 51},
        {"line": 101},
    ]
    assert "\n".join(_texts(result)).splitlines() == [
        f"row{i},value" for i in range(1, 121)
    ]


@pytest.mark.parametrize("data", [b"caf\xe9 latin-1", b"has a \x00 nul"])
def test_text_outside_the_adr_text_rule_fails(data):
    assert extract(data, ae.TEXT).status == ae.STATUS_FAILED


def test_corrupt_and_encrypted_pdfs_fail():
    assert extract(b"%PDF-1.4 not really a pdf", ae.PDF).status == ae.STATUS_FAILED
    writer = PdfWriter()
    writer.add_blank_page(612, 792)
    writer.encrypt(user_password="secret", algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    result = extract(out.getvalue(), ae.PDF)
    assert result.status == ae.STATUS_FAILED
    assert result.detail == "encrypted"


# --- G1: subprocess with a wall-clock timeout --------------------------------


def test_g1_a_hung_worker_is_killed_and_reported_as_timeout():
    hang = [sys.executable, "-c", "import time; time.sleep(60)"]
    started = time.monotonic()
    result = ae._run_worker(b"x", ae.TEXT, ExtractionLimits(timeout_seconds=0.5), hang)
    assert result == ae.ExtractionResult(ae.STATUS_TIMEOUT, detail="wall_clock")
    assert time.monotonic() - started < 10


def test_g1_extract_runs_in_a_child_bounded_by_the_wall_clock():
    # Interpreter start-up alone outlasts this limit, so the child is always
    # killed: a run in the calling process could not time out at all.
    result = extract(b"plain text", ae.TEXT, ExtractionLimits(timeout_seconds=0.001))
    assert result.status == ae.STATUS_TIMEOUT


def test_g1_a_crashing_worker_is_failed_not_raised():
    crash = [sys.executable, "-c", "import os; os._exit(3)"]
    result = ae._run_worker(b"x", ae.TEXT, ExtractionLimits(), crash)
    assert result.status == ae.STATUS_FAILED


def test_g1_worker_output_over_the_char_cap_is_rejected():
    flood = [
        sys.executable,
        "-c",
        'print(\'{"status":"ok","segments":[{"text":"xxxx","locator":{}}],"truncated":false}\')',
    ]
    result = ae._run_worker(b"x", ae.TEXT, ExtractionLimits(max_chars=3), flood)
    assert result.status == ae.STATUS_FAILED


def test_input_over_the_byte_cap_is_refused_without_starting_a_worker(monkeypatch):
    monkeypatch.setattr(ae, "_run_worker", lambda *a: pytest.fail("worker started"))
    result = extract(b"x" * 11, ae.TEXT, ExtractionLimits(max_input_bytes=10))
    assert result.status == ae.STATUS_TOO_LARGE


# --- G2: zip bombs are refused before decompression --------------------------

_ZIP_TYPES = [ae.DOCX, ae.XLSX, ae.PPTX, ae.ODT, ae.ODS, ae.ODP]


def _deflate_bomb(size: int = 20 * 1024 * 1024) -> bytes:
    return _zip(
        {"mimetype": "x", "[Content_Types].xml": "<Types/>", "bomb.xml": b"\0" * size}
    )


@pytest.mark.parametrize("ctype", _ZIP_TYPES)
def test_g2_high_ratio_member_is_refused_for_every_zip_type(ctype):
    result = extract(_deflate_bomb(), ctype)
    assert result.status == ae.STATUS_TOO_LARGE
    assert result.detail == "zip_ratio"


def _with_declared_size(data: bytes, member: str, size: int) -> bytes:
    """Rewrite a member's declared uncompressed size in both headers and wreck its data."""
    buf = bytearray(data)
    name = member.encode()
    local = buf.find(b"PK\x03\x04" + b"")
    while local != -1:
        name_len = struct.unpack_from("<H", buf, local + 26)[0]
        extra_len = struct.unpack_from("<H", buf, local + 28)[0]
        if bytes(buf[local + 30 : local + 30 + name_len]) == name:
            struct.pack_into("<I", buf, local + 22, size)
            data_start = local + 30 + name_len + extra_len
            compressed = struct.unpack_from("<I", buf, local + 18)[0]
            buf[data_start : data_start + compressed] = b"\xff" * compressed
        local = buf.find(b"PK\x03\x04", local + 4)
    central = buf.find(b"PK\x01\x02")
    while central != -1:
        name_len = struct.unpack_from("<H", buf, central + 28)[0]
        if bytes(buf[central + 46 : central + 46 + name_len]) == name:
            struct.pack_into("<I", buf, central + 24, size)
        central = buf.find(b"PK\x01\x02", central + 4)
    return bytes(buf)


def test_g2_declared_total_is_refused_before_any_member_is_inflated():
    # The member's data is garbage: inflating it would fail as a corrupt file.
    # too_large proves the declared size was judged first.
    data = _with_declared_size(_docx(), "word/document.xml", 0x7FFFFFFF)
    result = extract(data, ae.DOCX)
    assert result.status == ae.STATUS_TOO_LARGE
    assert result.detail == "zip_uncompressed_bytes"


def test_g2_declared_total_cap_is_configurable():
    data = _docx()
    assert extract(data, ae.DOCX).status == ae.STATUS_OK
    limits = ExtractionLimits(max_zip_uncompressed_bytes=1000)
    assert extract(data, ae.DOCX, limits).status == ae.STATUS_TOO_LARGE


def test_g2_member_count_is_capped():
    data = _zip({f"part{i}.xml": "<a/>" for i in range(30)})
    result = extract(data, ae.ODT, ExtractionLimits(max_zip_members=10))
    assert result.status == ae.STATUS_TOO_LARGE
    assert result.detail == "zip_members"


def test_g2_overlapping_entries_count_against_the_declared_total():
    # Overlapping-entry bombs reuse one compressed stream for many members; the
    # stream is stored once but each member's declared size counts.
    data = _zip({f"m{i}.xml": b"\0" * 400_000 for i in range(5)}, zipfile.ZIP_STORED)
    limits = ExtractionLimits(max_zip_uncompressed_bytes=1_000_000)
    assert extract(data, ae.ODT, limits).status == ae.STATUS_TOO_LARGE


def test_g2_small_highly_compressible_parts_are_not_mistaken_for_bombs():
    paragraphs = ["same words " * 50] * 200  # ~110 KB of XML that compresses very well
    result = extract(_docx(paragraphs=paragraphs), ae.DOCX)
    assert result.status == ae.STATUS_OK


# --- G3: no DTD or external entity is resolved -------------------------------


@pytest.fixture
def secret_file(tmp_path):
    path = tmp_path / "secret.txt"
    path.write_text("TOP-SECRET-MARKER")
    return path


def _xxe_doctype(path) -> str:
    return f'<!DOCTYPE x [<!ENTITY xxe SYSTEM "file://{path}">]>'


def _assert_refused_without_leak(result):
    assert result.status == ae.STATUS_FAILED
    assert result.detail == "xml_dtd"
    assert "TOP-SECRET-MARKER" not in "".join(_texts(result))


def test_g3_xlsx_external_entity_is_not_resolved(secret_file):
    data = _xlsx(
        {"S": [["s:1"]]}, doctype=_xxe_doctype(secret_file), entity_ref="&xxe;"
    )
    _assert_refused_without_leak(extract(data, ae.XLSX))


def test_g3_pptx_external_entity_is_not_resolved(secret_file):
    data = _pptx([["&xxe;"]], doctype=_xxe_doctype(secret_file))
    _assert_refused_without_leak(extract(data, ae.PPTX))


@pytest.mark.parametrize(
    "ctype,body", [(ae.ODT, "<office:text><text:p>&xxe;</text:p></office:text>")]
)
def test_g3_odf_external_entity_is_not_resolved(secret_file, ctype, body):
    data = _odf(ctype, body, doctype=_xxe_doctype(secret_file))
    _assert_refused_without_leak(extract(data, ctype))


def _replace_member(data: bytes, name: str, body: str) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(
                info.filename, body if info.filename == name else source.read(info)
            )
    return out.getvalue()


def test_g3_docx_external_entity_is_not_resolved(secret_file):
    document_xml = (
        zipfile.ZipFile(io.BytesIO(_docx(["PLACEHOLDER"])))
        .read("word/document.xml")
        .decode()
    )
    head, root = document_xml.split("?>", 1)
    poisoned = (
        head + "?>" + _xxe_doctype(secret_file) + root.replace("PLACEHOLDER", "&xxe;")
    )
    data = _replace_member(_docx(["PLACEHOLDER"]), "word/document.xml", poisoned)
    _assert_refused_without_leak(extract(data, ae.DOCX))


def test_g3_docx_dtd_in_a_part_python_docx_does_not_read_is_still_refused():
    dtd = '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "b">]><x>&a;</x>'
    data = _replace_member(_docx(), "docProps/app.xml", dtd)
    assert extract(data, ae.DOCX).detail == "xml_dtd"


def test_g3_billion_laughs_is_refused_quickly():
    entities = '<!ENTITY a0 "lol">' + "".join(
        f'<!ENTITY a{i} "{("&a" + str(i - 1) + ";") * 10}">' for i in range(1, 10)
    )
    data = _odf(
        ae.ODT,
        "<office:text><text:p>&a9;</text:p></office:text>",
        f"<!DOCTYPE x [{entities}]>",
    )
    started = time.monotonic()
    result = extract(data, ae.ODT)
    assert result.status == ae.STATUS_FAILED
    assert time.monotonic() - started < 10


def test_g3_external_relationship_targets_are_never_followed():
    data = _xlsx({"Sheet": [["1"]]})
    rels = _rels([("rId1", _REL_WS, "file:///etc/passwd")]).replace(
        "/>", ' TargetMode="External"/>', 1
    )
    result = extract(_replace_member(data, "xl/_rels/workbook.xml.rels", rels), ae.XLSX)
    assert result.status == ae.STATUS_OK
    assert result.segments == ()


# --- G4: per-file character cap ----------------------------------------------


def test_g4_text_is_cut_at_the_cap_and_flagged_truncated():
    result = extract(
        ("word " * 2000).encode(), ae.TEXT, ExtractionLimits(max_chars=1000)
    )
    assert result.status == ae.STATUS_OK
    assert result.truncated is True
    assert _total_chars(result) == 1000


def test_g4_cap_spans_segments_across_sheets():
    data = _xlsx({"A": [["i:" + "a" * 600]], "B": [["i:" + "b" * 600]], "C": [["i:c"]]})
    result = extract(data, ae.XLSX, ExtractionLimits(max_chars=1000))
    assert result.truncated is True
    assert _total_chars(result) == 1000
    assert [s.locator for s in result.segments] == [{"sheet": "A"}, {"sheet": "B"}]


def test_g4_cap_applies_to_pdf_pages():
    result = extract(
        _pdf(["A" * 80, "B" * 80, "C" * 80]), ae.PDF, ExtractionLimits(max_chars=100)
    )
    assert result.truncated is True
    assert _total_chars(result) == 100


def test_g4_text_exactly_at_the_cap_is_not_truncated():
    result = extract(b"x" * 50, ae.TEXT, ExtractionLimits(max_chars=50))
    assert result.truncated is False
    assert _total_chars(result) == 50


def test_g4_collector_never_exceeds_its_budget():
    out = ae._Collector(10, max_segments=100)
    assert out.add("12345678", {}) is True
    assert out.add("abcdef", {}) is False
    assert out.add("more", {}) is False
    assert sum(len(s.text) for s in out.segments) == 10
    assert out.truncated is True


# --- G5: image-only PDF ------------------------------------------------------


def test_g5_pdf_without_a_text_layer_is_no_text_layer_not_empty_ok():
    result = extract(_blank_pdf(), ae.PDF)
    assert result.status == ae.STATUS_NO_TEXT_LAYER
    assert result.segments == ()


def test_g5_a_pdf_with_some_text_is_ok_and_skips_empty_pages():
    result = extract(_pdf(["", "Only page two"]), ae.PDF)
    assert result.status == ae.STATUS_OK
    assert [s.locator for s in result.segments] == [{"page": 2}]


# --- G6: dispatch by sniffed content type only -------------------------------


def test_g6_extract_takes_no_file_name():
    assert list(inspect.signature(extract).parameters) == [
        "data",
        "content_type",
        "limits",
    ]


def test_g6_pdf_bytes_labelled_text_are_read_as_text_not_parsed_as_pdf():
    result = extract(_pdf(["Hidden"]), ae.TEXT)
    assert result.status == ae.STATUS_OK
    assert result.segments[0].text.startswith("%PDF-1.4")


def test_g6_docx_bytes_labelled_xlsx_are_parsed_as_xlsx():
    result = extract(_docx(), ae.XLSX)
    assert result.status == ae.STATUS_FAILED
    assert result.detail == "missing_part"


@pytest.mark.parametrize(
    "ctype",
    ["image/png", "image/jpeg", "text/html", "application/zip", "application/xml", ""],
)
def test_g6_types_without_an_extractor_are_unsupported(ctype):
    assert extract(b"anything", ctype).status == ae.STATUS_UNSUPPORTED_TYPE


def test_g6_content_type_parameters_and_case_are_ignored():
    result = extract(b"hello", "Text/Plain; charset=utf-8")
    assert result.status == ae.STATUS_OK
    assert _texts(result) == ["hello"]


# --- Review regressions ------------------------------------------------------


def _fake_worker(code: str) -> list[str]:
    return [sys.executable, "-c", code]


def test_g1_worker_runs_without_the_parent_environment(monkeypatch):
    monkeypatch.setenv("ATTACHMENT_EXTRACT_TEST_SECRET", "s3cret")
    code = (
        "import json, os; print(json.dumps({'status': 'ok', 'truncated': False, "
        "'segments': [{'text': os.environ.get('ATTACHMENT_EXTRACT_TEST_SECRET', 'absent'), "
        "'locator': {}}]}))"
    )
    result = ae._run_worker(b"x", ae.TEXT, ExtractionLimits(), _fake_worker(code))
    assert _texts(result) == ["absent"]


def test_g1_oversized_or_hostile_worker_replies_are_failed_not_raised():
    flood = _fake_worker("import sys; sys.stdout.write('x' * 5000)")
    limits = ExtractionLimits(max_chars=10, max_segments=1)
    assert ae._run_worker(b"x", ae.TEXT, limits, flood).status == ae.STATUS_FAILED
    nested = _fake_worker("import sys; sys.stdout.write('[' * 200000)")
    assert (
        ae._run_worker(b"x", ae.TEXT, ExtractionLimits(), nested).status
        == ae.STATUS_FAILED
    )


@pytest.mark.parametrize(
    "segments",
    [
        [{"text": "a", "locator": {}}] * 3,
        [{"text": "a", "locator": {"sheet": "n" * 300}}],
        [{"text": "a", "locator": {"evil": 1}}],
    ],
)
def test_g1_parent_rejects_replies_over_segment_or_locator_bounds(segments):
    limits = ExtractionLimits(max_segments=2)
    obj = {"status": "ok", "segments": segments, "truncated": False}
    with pytest.raises((ValueError, TypeError)):
        ae._from_wire(obj, limits)


def _under_declare(data: bytes, member: str, declared: bytes) -> bytes:
    """Make ``member`` declare ``declared``'s size and CRC while its stream holds more."""
    buf = bytearray(data)
    name = member.encode()
    crc = zlib.crc32(declared)
    for signature, name_len_at, name_at, crc_at, size_at in (
        (b"PK\x03\x04", 26, 30, 14, 22),
        (b"PK\x01\x02", 28, 46, 16, 24),
    ):
        at = buf.find(signature)
        while at != -1:
            name_len = struct.unpack_from("<H", buf, at + name_len_at)[0]
            if bytes(buf[at + name_at : at + name_at + name_len]) == name:
                struct.pack_into("<I", buf, at + crc_at, crc)
                struct.pack_into("<I", buf, at + size_at, len(declared))
            at = buf.find(signature, at + 4)
    return bytes(buf)


def test_g2_a_member_that_under_declares_its_size_is_never_inflated_past_it():
    document_xml = zipfile.ZipFile(io.BytesIO(_docx(["Bounded"]))).read(
        "word/document.xml"
    )
    source = zipfile.ZipFile(io.BytesIO(_docx(["Bounded"])))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename != "word/document.xml":
                target.writestr(info.filename, source.read(info))
                continue
            with target.open("word/document.xml", "w") as member:
                member.write(document_xml)
                for _ in range(64):
                    member.write(b" " * (1024 * 1024))
    data = _under_declare(out.getvalue(), "word/document.xml", document_xml)

    tracemalloc.start()
    try:
        result = ae._extract_in_process(data, ae.DOCX, ExtractionLimits())
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert result.status == ae.STATUS_OK
    assert _texts(result) == ["Bounded"]
    assert peak < 16 * 1024 * 1024


def test_g3_docx_dtd_in_a_part_not_named_xml_is_refused():
    data = _docx(["PLACEHOLDER"])
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            body = source.read(info).decode("utf-8", "surrogateescape")
            name = info.filename
            if name == "word/document.xml":
                name = "word/document.dat"
                head, root = body.split("?>", 1)
                body = head + '?><!DOCTYPE w:document [<!ENTITY x "EXPANDED">]>' + root
                body = body.replace("PLACEHOLDER", "before &x;")
            else:
                body = body.replace("document.xml", "document.dat")
            target.writestr(name, body.encode("utf-8", "surrogateescape"))
    result = extract(out.getvalue(), ae.DOCX)
    assert result.status == ae.STATUS_FAILED
    assert result.detail == "xml_dtd"


def test_g4_a_cut_on_a_row_boundary_of_the_last_segment_is_reported_truncated():
    result = extract(
        _xlsx({"S": [["i:abcde"], ["i:fgh"]]}), ae.XLSX, ExtractionLimits(max_chars=5)
    )
    assert _texts(result) == ["abcde"]
    assert result.truncated is True


def test_g4_bounded_join_stops_only_past_the_limit():
    assert ae._bounded_join(iter(["abcde", "fgh"]), "\n", 5) == "abcde\nfgh"
    assert ae._bounded_join(iter(["ab", "cd"]), "\n", 5) == "ab\ncd"


def test_g4_segment_count_is_capped_and_reported_truncated():
    result = extract(b"a\n\n" * 30, ae.TEXT, ExtractionLimits(max_segments=10))
    assert len(result.segments) == 10
    assert result.truncated is True


def test_g4_locator_strings_are_capped_and_repeated_sheets_read_once():
    data = _xlsx({"N" * 1000: [["i:x"]]})
    source = zipfile.ZipFile(io.BytesIO(data))
    workbook = source.read("xl/workbook.xml").decode()
    repeated = workbook.replace(
        "</sheets>", '<sheet name="again" sheetId="9" r:id="rId1"/></sheets>'
    )
    result = extract(_replace_member(data, "xl/workbook.xml", repeated), ae.XLSX)
    assert result.segments == (ae.Segment("x", {"sheet": "N" * 64}),)


def _docx_with_body(inner: str) -> bytes:
    data = _docx(["PLACEHOLDER"])
    document_xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
    start = (
        document_xml.index("<w:p>")
        if "<w:p>" in document_xml
        else document_xml.index("<w:p ")
    )
    end = document_xml.index("</w:p>", start) + len("</w:p>")
    return _replace_member(
        data, "word/document.xml", document_xml[:start] + inner + document_xml[end:]
    )


def test_g7_docx_body_level_content_controls_are_read():
    inner = "<w:sdt><w:sdtPr/><w:sdtContent><w:p><w:r><w:t>INSIDE_SDT</w:t></w:r></w:p></w:sdtContent></w:sdt>"
    result = extract(_docx_with_body(inner), ae.DOCX)
    assert result.segments == (ae.Segment("INSIDE_SDT", {"paragraph": 1}),)


def test_g7_docx_nested_table_text_is_kept():
    def table(cell: str) -> str:
        return f"<w:tbl><w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc>{cell}</w:tc></w:tr></w:tbl>"

    inner = table(table("<w:p><w:r><w:t>NESTED</w:t></w:r></w:p>") + "<w:p/>")
    result = extract(_docx_with_body(inner), ae.DOCX)
    assert result.segments == (ae.Segment("NESTED", {"table": 1}),)


def test_g7_text_lines_are_counted_by_newline_only():
    result = extract("a\x0cb\r\n\r\nc\n".encode(), ae.TEXT)
    assert result.segments == (
        ae.Segment("a\x0cb", {"line": 1}),
        ae.Segment("c", {"line": 3}),
    )


def test_g7_pptx_line_breaks_separate_runs():
    data = _pptx([["Line1</a:t></a:r><a:br/><a:r><a:t>Line2"]])
    assert _texts(extract(data, ae.PPTX)) == ["Line1\nLine2"]


def test_g7_odt_comment_metadata_is_not_spliced_into_paragraph_text():
    body = (
        '<office:text><text:p>Body<office:annotation xmlns:dc="http://purl.org/dc/elements/1.1/">'
        "<dc:creator>Reviewer</dc:creator><dc:date>2026-01-01</dc:date>"
        "<text:p>Comment</text:p></office:annotation> text</text:p></office:text>"
    )
    assert _texts(extract(_odf(ae.ODT, body), ae.ODT)) == ["Body text", "Comment"]


# --- Mutation-review survivors closed ----------------------------------------


def test_g1_a_worker_exiting_non_zero_is_failed_even_with_valid_output():
    code = (
        'import sys; sys.stdout.write(\'{"status": "ok", "segments": [], '
        '"truncated": false}\'); sys.exit(1)'
    )
    result = ae._run_worker(b"x", ae.TEXT, ExtractionLimits(), _fake_worker(code))
    assert result == ae.ExtractionResult(ae.STATUS_FAILED, detail="worker_exit_1")


def test_g1_an_unknown_worker_status_is_failed():
    code = 'print(\'{"status": "bogus", "segments": [], "truncated": false}\')'
    result = ae._run_worker(b"x", ae.TEXT, ExtractionLimits(), _fake_worker(code))
    assert result == ae.ExtractionResult(ae.STATUS_FAILED, detail="worker_output")


def test_g1_the_child_runs_under_its_address_space_limit():
    data = b"word " * (10 * 1024 * 1024)  # 50 MiB of text, decoded to a larger str
    limits = ExtractionLimits(
        max_input_bytes=len(data), memory_limit_bytes=100 * 1024 * 1024
    )
    result = extract(data, ae.TEXT, limits)
    assert result.status == ae.STATUS_FAILED


def test_g2_total_ratio_is_checked_when_every_member_is_under_the_floor():
    data = _zip({f"m{i}.xml": b"\0" * 600_000 for i in range(5)})
    result = extract(data, ae.ODT)
    assert (result.status, result.detail) == (ae.STATUS_TOO_LARGE, "zip_ratio")


def test_g2_total_ratio_has_no_size_floor():
    data = _zip({"one": b"a" * 250_000, "two": b"a" * 250_000})
    assert 500_000 < ExtractionLimits().zip_ratio_floor_bytes
    assert 500_000 / len(data) > ExtractionLimits().max_zip_ratio
    with pytest.raises(ae._TooLarge, match="zip_ratio"):
        ae._open_zip(data, ExtractionLimits())


def test_g2_member_ratio_is_checked_when_the_total_ratio_is_low():
    import os

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("bomb.xml", b"\0" * 2_000_000, zipfile.ZIP_DEFLATED)
        archive.writestr("noise.bin", os.urandom(3_000_000), zipfile.ZIP_STORED)
    result = extract(out.getvalue(), ae.ODT)
    assert (result.status, result.detail) == (ae.STATUS_TOO_LARGE, "zip_ratio")


def test_g3_docx_dtd_in_a_relationships_part_is_refused():
    rels = _rels([]).replace("?>", '?><!DOCTYPE x [<!ENTITY a "b">]>', 1)
    data = _replace_member(_docx(), "word/_rels/document.xml.rels", rels)
    assert extract(data, ae.DOCX).detail == "xml_dtd"


def test_g3_xml_nested_past_the_parser_depth_limit_is_refused():
    deep = "<a>" * 300 + "</a>" * 300
    body = f"<office:text><text:p>{deep}</text:p></office:text>"
    assert extract(_odf(ae.ODT, body), ae.ODT).status == ae.STATUS_FAILED


def test_g4_bounded_join_terminates_on_endless_parts():
    import itertools

    assert len(ae._bounded_join(itertools.repeat("x"), "\n", 100)) == 101


def test_g4_a_single_sheet_far_over_the_cap_is_cut_exactly():
    rows = [[f"i:row{i:05d}"] for i in range(2000)]
    result = extract(_xlsx({"Big": rows}), ae.XLSX, ExtractionLimits(max_chars=1000))
    assert _total_chars(result) == 1000
    assert result.truncated is True


def test_g7_odf_space_runs_are_capped():
    body = '<office:text><text:p>a<text:s text:c="900000000"/>b</text:p></office:text>'
    result = extract(_odf(ae.ODT, body), ae.ODT)
    assert _texts(result) == ["a" + " " * 100 + "b"]


# --- Review round 2 regressions ----------------------------------------------


@pytest.mark.parametrize(
    "wrapper",
    [
        "<w:ins>{}</w:ins>",
        "<w:sdt><w:sdtPr/><w:sdtContent>{}</w:sdtContent></w:sdt>",
        '<w:fldSimple w:instr="AUTHOR">{}</w:fldSimple>',
        "<w:smartTag>{}</w:smartTag>",
    ],
)
def test_g7_docx_runs_inside_inline_wrappers_are_read(wrapper):
    run = "<w:r><w:t>WRAPPED</w:t></w:r>"
    inner = f'<w:p><w:r><w:t xml:space="preserve">before </w:t></w:r>{wrapper.format(run)}</w:p>'
    assert _texts(extract(_docx_with_body(inner), ae.DOCX)) == ["before WRAPPED"]


def test_g7_docx_tracked_deletions_and_field_codes_are_left_out():
    inner = (
        "<w:p><w:r><w:t>kept</w:t></w:r>"
        "<w:del><w:r><w:delText>DELETED</w:delText></w:r></w:del>"
        '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
        "<w:r><w:tab/><w:t>after</w:t></w:r></w:p>"
    )
    assert _texts(extract(_docx_with_body(inner), ae.DOCX)) == ["kept\tafter"]


def test_g7_odt_tracked_deletions_are_not_paragraphs():
    body = (
        "<office:text><text:tracked-changes><text:changed-region><text:deletion>"
        "<text:p>DELETED_TEXT</text:p></text:deletion></text:changed-region></text:tracked-changes>"
        "<text:p>First</text:p><text:p>Second</text:p></office:text>"
    )
    result = extract(_odf(ae.ODT, body), ae.ODT)
    assert result.segments == (
        ae.Segment("First", {"paragraph": 1}),
        ae.Segment("Second", {"paragraph": 2}),
    )


def test_g7_pptx_alternate_content_is_read_once():
    _mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    choice = "<a:p><a:r><a:t>EQUATION</a:t></a:r></a:p>"
    fallback = "<a:p><a:r><a:t>FALLBACK_COPY</a:t></a:r></a:p>"
    data = _pptx([["Title"]])
    slide = zipfile.ZipFile(io.BytesIO(data)).read("ppt/slides/slide1.xml").decode()
    alternate = (
        f'<mc:AlternateContent xmlns:mc="{_mc}"><mc:Choice Requires="a14"><p:sp><p:txBody>'
        f"{choice}</p:txBody></p:sp></mc:Choice><mc:Fallback><p:sp><p:txBody>{fallback}"
        "</p:txBody></p:sp></mc:Fallback></mc:AlternateContent>"
    )
    slide = slide.replace("</p:spTree>", alternate + "</p:spTree>")
    result = extract(_replace_member(data, "ppt/slides/slide1.xml", slide), ae.PPTX)
    assert _texts(result) == ["Title\nEQUATION"]


def test_g7_pptx_a_slide_listed_twice_is_read_once():
    data = _pptx([["Only"]])
    presentation = (
        zipfile.ZipFile(io.BytesIO(data)).read("ppt/presentation.xml").decode()
    )
    presentation = presentation.replace(
        "</p:sldIdLst>", '<p:sldId id="900" r:id="rId1"/></p:sldIdLst>'
    )
    result = extract(
        _replace_member(data, "ppt/presentation.xml", presentation), ae.PPTX
    )
    assert result.segments == (ae.Segment("Only", {"slide": 1}),)


def test_g7_docx_with_an_svg_image_carrying_a_doctype_still_extracts():
    data = _docx(["Has a picture"])
    types = zipfile.ZipFile(io.BytesIO(data)).read("[Content_Types].xml").decode()
    types = types.replace(
        "</Types>", '<Default Extension="svg" ContentType="image/svg+xml"/></Types>'
    )
    data = _replace_member(data, "[Content_Types].xml", types)
    svg = (
        '<?xml version="1.0"?><!DOCTYPE svg PUBLIC "-//W3C//DTD SVG 1.1//EN" '
        '"http://www.w3.org/Graphics/SVG/1.1/DTD/svg11.dtd"><svg/>'
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        source = zipfile.ZipFile(io.BytesIO(data))
        for info in source.infolist():
            target.writestr(info.filename, source.read(info))
        target.writestr("word/media/image1.svg", svg)
    assert _texts(extract(out.getvalue(), ae.DOCX)) == ["Has a picture"]


def test_g3_docx_relationships_part_is_scanned_by_name_without_a_content_type():
    data = _docx()
    types = zipfile.ZipFile(io.BytesIO(data)).read("[Content_Types].xml").decode()
    types = types.replace('Extension="rels"', 'Extension="relsx"')
    data = _replace_member(data, "[Content_Types].xml", types)
    rels = _rels([]).replace("?>", '?><!DOCTYPE x [<!ENTITY a "b">]>', 1)
    data = _replace_member(data, "word/_rels/document.xml.rels", rels)
    assert extract(data, ae.DOCX).detail == "xml_dtd"


@pytest.mark.parametrize("method", [zipfile.ZIP_BZIP2, zipfile.ZIP_LZMA])
def test_g2_compression_methods_outside_ooxml_and_odf_are_refused_unread(method):
    data = _zip({"mimetype": ae.ODT, "content.xml": "<a/>"}, method)
    result = extract(data, ae.ODT)
    assert (result.status, result.detail) == (ae.STATUS_FAILED, "zip_compression")


def test_g1_a_valid_reply_over_the_byte_cap_is_failed():
    code = (
        'import sys; sys.stdout.write(\'{"status": "ok", "segments": [], '
        "\"truncated\": false}' + ' ' * 100000)"
    )
    limits = ExtractionLimits(max_chars=10, max_segments=1)
    result = ae._run_worker(b"x", ae.TEXT, limits, _fake_worker(code))
    assert result == ae.ExtractionResult(ae.STATUS_FAILED, detail="worker_output")


@pytest.mark.parametrize(
    "locator", [{"paragraph": "\x01" * 64}, {"paragraph": 2**53 - 1}]
)
def test_g1_a_reply_at_the_locator_and_segment_bounds_fits_the_byte_cap(locator):
    import json

    limits = ExtractionLimits(max_chars=1000, max_segments=1000)
    segment = ae.Segment("\x01", locator)
    reply = ae._to_wire(ae.ExtractionResult("ok", (segment,) * 1000, False, "x"))
    assert ae._from_wire(reply, limits).segments[0] == segment
    encoded = json.dumps(reply, ensure_ascii=False).encode()
    assert len(encoded) <= ae._max_reply_bytes(limits)


def test_g1_reply_byte_cap_matches_the_adr_formula():
    # ADR 0010: 6 x chars + 512 x segments, plus constant reply framing.
    assert ae._max_reply_bytes(ExtractionLimits()) == 11_124_096


def test_g1_worker_argv_isolates_the_child_interpreter():
    argv = ae._worker_argv()
    assert argv[0] == sys.executable and argv[1] == "-I"


def test_g4_collector_replaces_lone_surrogates():
    out = ae._Collector(100, max_segments=10)
    out.add("a\ud800b", {})
    out.segments[0].text.encode("utf-8")


# --- Review round 2 mutation findings ----------------------------------------


def test_g1_a_child_flooding_stdout_is_stopped_at_the_reply_cap():
    flood = _fake_worker(
        "import sys\nwhile True: sys.stdout.buffer.write(b'x' * 65536)"
    )
    limits = ExtractionLimits(max_chars=10, max_segments=1, timeout_seconds=20)
    started = time.monotonic()
    result = ae._run_worker(b"x", ae.TEXT, limits, flood)
    assert result == ae.ExtractionResult(ae.STATUS_FAILED, detail="worker_output")
    # Stopped at the cap, not at the wall clock.
    assert time.monotonic() - started < 10


@pytest.mark.parametrize(
    "locator",
    [
        [],
        "x",
        {"page": True},
        {"page": 1.5},
        {"page": None},
        {"sheet": "n" * 65},
        {"page": -1},
        {"page": 2**53},
        {"page": 1, "sheet": "a"},
    ],
)
def test_g1_parent_rejects_malformed_locators(locator):
    obj = {
        "status": "ok",
        "segments": [{"text": "a", "locator": locator}],
        "truncated": False,
    }
    with pytest.raises((ValueError, TypeError)):
        ae._from_wire(obj, ExtractionLimits())


@pytest.mark.parametrize(
    "ctype,builder,rels_name",
    [
        (ae.XLSX, lambda: _xlsx({"S": [["1"]]}), "xl/_rels/workbook.xml.rels"),
        (ae.PPTX, lambda: _pptx([["x"]]), "ppt/_rels/presentation.xml.rels"),
    ],
)
def test_g3_dtd_in_a_relationships_part_is_refused(ctype, builder, rels_name):
    data = builder()
    rels = zipfile.ZipFile(io.BytesIO(data)).read(rels_name).decode()
    rels = rels.replace("?>", '?><!DOCTYPE x [<!ENTITY a "b">]>', 1)
    assert extract(_replace_member(data, rels_name, rels), ctype).detail == "xml_dtd"


def test_g7_docx_nested_table_keeps_every_row():
    def table(rows: list[str]) -> str:
        body = "".join(f"<w:tr><w:tc>{cell}</w:tc></w:tr>" for cell in rows)
        return f"<w:tbl><w:tblGrid><w:gridCol/></w:tblGrid>{body}</w:tbl>"

    def para(text: str) -> str:
        return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"

    inner = table([table([para("N1"), para("N2")]) + "<w:p/>"])
    assert _texts(extract(_docx_with_body(inner), ae.DOCX)) == ["N1\nN2"]


def test_g5_a_pdf_whose_pages_hold_only_whitespace_is_no_text_layer():
    assert extract(_pdf([" ", "  "]), ae.PDF).status == ae.STATUS_NO_TEXT_LAYER


# --- Review round 3 regressions ----------------------------------------------


def test_g3_docx_relationships_part_typed_as_an_image_is_still_scanned():
    data = _docx()
    types = zipfile.ZipFile(io.BytesIO(data)).read("[Content_Types].xml").decode()
    types = types.replace(
        "</Types>",
        '<Override PartName="/word/_rels/document.xml.rels" ContentType="image/png"/></Types>',
    )
    data = _replace_member(data, "[Content_Types].xml", types)
    rels = (
        zipfile.ZipFile(io.BytesIO(data)).read("word/_rels/document.xml.rels").decode()
    )
    rels = rels.replace("?>", '?><!DOCTYPE Relationships [<!ENTITY x "boom">]>', 1)
    data = _replace_member(data, "word/_rels/document.xml.rels", rels)
    assert extract(data, ae.DOCX).detail == "xml_dtd"


def test_g7_docx_non_breaking_hyphens_and_positional_tabs_are_kept():
    inner = (
        "<w:p><w:r><w:t>non</w:t><w:noBreakHyphen/><w:t>breaking</w:t>"
        '<w:ptab w:relativeTo="margin" w:alignment="right" w:leader="none"/><w:t>end</w:t></w:r></w:p>'
    )
    assert _texts(extract(_docx_with_body(inner), ae.DOCX)) == ["non-breaking\tend"]


@pytest.mark.skipif(sys.platform == "win32", reason="process groups are POSIX")
def test_g1_a_process_the_child_leaves_behind_does_not_hold_the_call_open():
    code = (
        "import subprocess, sys\n"
        "subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(30)'])\n"
        "sys.exit(0)"
    )
    started = time.monotonic()
    result = ae._run_worker(
        b"x", ae.TEXT, ExtractionLimits(timeout_seconds=3), _fake_worker(code)
    )
    assert result.status == ae.STATUS_FAILED
    assert time.monotonic() - started < 3


def test_g7_docx_breaks_text_boxes_and_alternate_content_read_as_current_text():
    _mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    inner = (
        "<w:p><w:r><w:t>a</w:t><w:br/><w:t>b</w:t><w:cr/><w:t>c</w:t></w:r>"
        f'<w:r><mc:AlternateContent xmlns:mc="{_mc}"><mc:Choice Requires="wps">'
        "<w:p><w:r><w:t>BOX</w:t></w:r></w:p></mc:Choice>"
        "<mc:Fallback><w:p><w:r><w:t>FALLBACK_COPY</w:t></w:r></w:p></mc:Fallback>"
        "</mc:AlternateContent></w:r><w:r><w:t>after</w:t></w:r>"
        "<w:moveFrom><w:r><w:t>MOVED_AWAY</w:t></w:r></w:moveFrom></w:p>"
    )
    assert _texts(extract(_docx_with_body(inner), ae.DOCX)) == ["a\nb\nc\nBOX\nafter"]


def test_g7_docx_table_cells_read_tracked_insertions():
    cell = "<w:p><w:r><w:t>cell </w:t></w:r><w:ins><w:r><w:t>INSERTED</w:t></w:r></w:ins></w:p>"
    inner = f"<w:tbl><w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc>{cell}</w:tc></w:tr></w:tbl>"
    assert _texts(extract(_docx_with_body(inner), ae.DOCX)) == ["cell INSERTED"]


# --- Review round 4 regressions ----------------------------------------------


@pytest.mark.skipif(sys.platform == "win32", reason="process groups are POSIX")
@pytest.mark.parametrize("held", ["stdout", "stdin"])
def test_g1_a_process_that_escapes_the_group_cannot_hold_parent_resources(held):
    import os
    import threading

    close_stdout = "import os; os.close(1); " if held == "stdin" else ""
    code = (
        "import subprocess, sys\n"
        "subprocess.Popen([sys.executable, '-c', "
        f"'{close_stdout}import time; time.sleep(4)'], start_new_session=True)\n"
        "sys.exit(0)"
    )
    fds_before = (
        len(os.listdir("/proc/self/fd")) if os.path.isdir("/proc/self/fd") else None
    )
    threads_before = threading.active_count()
    started = time.monotonic()
    result = ae._run_worker(
        b"x" * (4 * 1024 * 1024),
        ae.TEXT,
        ExtractionLimits(timeout_seconds=3),
        _fake_worker(code),
    )
    assert result.status in (ae.STATUS_FAILED, ae.STATUS_TIMEOUT)
    # Well inside the 5 s a blocked join would cost on top of the wall clock.
    assert time.monotonic() - started < 6
    assert threading.active_count() == threads_before
    if fds_before is not None:
        assert len(os.listdir("/proc/self/fd")) == fds_before


def test_g1_a_worker_that_cannot_be_started_is_failed_not_raised():
    result = ae._run_worker(
        b"x", ae.TEXT, ExtractionLimits(), ["/nonexistent/interpreter"]
    )
    assert result == ae.ExtractionResult(ae.STATUS_FAILED, detail="worker_spawn")


def test_g7_docx_horizontally_merged_cell_is_reported_once():
    data = _docx(paragraphs=[], table=[["MERGED", ""], ["a", "b"]])
    document = docx.Document(io.BytesIO(data))
    grid = document.tables[0]
    grid.cell(0, 0).merge(grid.cell(0, 1))
    grid.cell(0, 0).paragraphs[-1].text = ""
    out = io.BytesIO()
    document.save(out)
    result = extract(out.getvalue(), ae.DOCX)
    assert _texts(result) == ["MERGED\na\tb"]


# --- Review round 5 regressions ----------------------------------------------


@pytest.mark.skipif(sys.platform == "win32", reason="the selector exchange is POSIX")
def test_g1_a_reply_written_just_before_exit_is_read_in_full(monkeypatch):
    import selectors

    class SlowSelector(selectors.DefaultSelector):
        # Widens the gap between an empty select and the exit check, as a busy
        # server thread holding the GIL would.
        def select(self, timeout=None):
            events = super().select(timeout)
            if not events:
                time.sleep(0.3)
            return events

    monkeypatch.setattr(selectors, "DefaultSelector", SlowSelector)
    code = (
        "import sys, time; sys.stdin.buffer.read(); time.sleep(0.1); "
        'sys.stdout.write(\'{"status": "ok", "segments": [{"text": "tail", '
        '"locator": {}}], "truncated": false}\')'
    )
    result = ae._run_worker(
        b"x", ae.TEXT, ExtractionLimits(timeout_seconds=10), _fake_worker(code)
    )
    assert _texts(result) == ["tail"]


@pytest.mark.skipif(sys.platform == "win32", reason="the selector exchange is POSIX")
def test_g1_a_selector_that_cannot_be_created_is_failed_not_raised(monkeypatch):
    import selectors

    def no_descriptors():
        raise OSError(24, "Too many open files")

    monkeypatch.setattr(selectors, "DefaultSelector", no_descriptors)
    assert extract(b"hi", ae.TEXT) == ae.ExtractionResult(
        ae.STATUS_FAILED, detail="worker_io"
    )


_OK_REPLY = '{"status": "ok", "segments": [{"text": "done", "locator": {}}], "truncated": false}'


def test_g1_a_reply_whose_writer_closes_stdout_before_exiting_is_kept():
    code = (
        f"import os, sys, time; sys.stdout.write({_OK_REPLY!r}); sys.stdout.flush(); "
        "os.close(1); time.sleep(0.3)"
    )
    result = ae._run_worker(
        b"x", ae.TEXT, ExtractionLimits(timeout_seconds=10), _fake_worker(code)
    )
    assert _texts(result) == ["done"]


def test_g1_a_child_that_closes_stdout_and_hangs_times_out():
    code = "import os, time; os.close(1); time.sleep(30)"
    result = ae._run_worker(
        b"x", ae.TEXT, ExtractionLimits(timeout_seconds=1), _fake_worker(code)
    )
    assert result == ae.ExtractionResult(ae.STATUS_TIMEOUT, detail="wall_clock")


@pytest.mark.parametrize("code", ["pass", "import time; time.sleep(30)"])
def test_g1_a_payload_larger_than_the_pipe_is_never_blocking(code):
    started = time.monotonic()
    result = ae._run_worker(
        b"x" * (4 * 1024 * 1024),
        ae.TEXT,
        ExtractionLimits(timeout_seconds=1),
        _fake_worker(code),
    )
    assert result.status in (ae.STATUS_FAILED, ae.STATUS_TIMEOUT)
    assert time.monotonic() - started < 4


@pytest.mark.parametrize("extra", [0, 1, 4668])
def test_g1_reply_byte_cap_accepts_exactly_the_limit(extra):
    limits = ExtractionLimits(max_chars=10, max_segments=1)
    cap = ae._max_reply_bytes(limits)
    assert cap == 4668
    code = (
        f"import sys; reply = {_OK_REPLY!r}; "
        f"sys.stdout.write(reply + ' ' * ({cap + extra} - len(reply)))"
    )
    result = ae._run_worker(b"x", ae.TEXT, limits, _fake_worker(code))
    if extra:
        assert result == ae.ExtractionResult(ae.STATUS_FAILED, detail="worker_output")
    else:
        assert result.status == ae.STATUS_OK
        assert _texts(result) == ["done"]


@pytest.mark.parametrize(
    "field,detail",
    [
        ("max_zip_members", "zip_members"),
        ("max_zip_uncompressed_bytes", "zip_uncompressed_bytes"),
    ],
)
def test_g2_zip_declared_caps_accept_exactly_the_limit(field, detail):
    import os

    data = _zip({"one": os.urandom(250_000), "two": os.urandom(250_000)})
    cap = 2 if field == "max_zip_members" else 500_000
    with ae._open_zip(data, ExtractionLimits(**{field: cap})) as archive:
        assert len(archive.infolist()) == 2
    with pytest.raises(ae._TooLarge, match=detail):
        ae._open_zip(data, ExtractionLimits(**{field: cap - 1}))


def test_g3_xml_parser_disables_entities_dtd_loading_and_network(monkeypatch):
    from lxml import etree

    real_parser = etree.XMLParser
    seen = []

    def parser(**kwargs):
        seen.append(kwargs)
        return real_parser(**kwargs)

    monkeypatch.setattr(etree, "XMLParser", parser)
    assert ae._parse_xml(b"<root/>").tag == "root"
    assert len(seen) == 1
    for name, value in {
        "resolve_entities": False,
        "load_dtd": False,
        "no_network": True,
    }.items():
        assert seen[0][name] is value


@pytest.mark.parametrize(
    "name,entry",
    [
        (
            "word/extra.dat",
            '<Default Extension="dat" ContentType="application/custom+xml"/>',
        ),
        (
            "word/Extra.DaT",
            '<Default Extension="dAt" ContentType="application/custom+xml"/>',
        ),
        (
            "word/Extra.DaT",
            '<Override PartName="/WORD/extra.dAT" ContentType="application/custom+xml"/>',
        ),
        ("word/extra.XML", ""),
        ("word/_rels/extra.ReLs", ""),
    ],
)
def test_g3_docx_scans_default_override_and_untyped_xml_parts(name, entry):
    members = {
        "[Content_Types].xml": f"<Types>{entry}</Types>".encode(),
        name: b'<!DOCTYPE x [<!ENTITY e "value">]><x/>',
    }
    with pytest.raises(ae._Refused, match="xml_dtd"):
        ae._refuse_docx_dtds(members)


def test_g3_docx_missing_content_types_has_a_specific_detail():
    result = extract(_zip({"word/document.xml": "<document/>"}), ae.DOCX)
    assert result == ae.ExtractionResult(ae.STATUS_FAILED, detail="missing_part")


@pytest.mark.skipif(sys.platform == "win32", reason="resource limits require POSIX")
def test_g1_worker_applies_cpu_and_file_size_limits_before_parsing():
    import json

    # Run the real worker entry point, replacing only parsing with a probe.
    code = f"""
import json, resource, runpy
ns = runpy.run_path({ae.__file__!r})
def probe(*args):
    values = [resource.getrlimit(resource.RLIMIT_CPU), resource.getrlimit(resource.RLIMIT_FSIZE)]
    return ns['ExtractionResult']('ok', (ns['Segment'](json.dumps(values)),))
ns['_worker_main'].__globals__['_extract_in_process'] = probe
raise SystemExit(ns['_worker_main']())
"""
    result = ae._run_worker(
        b"x", ae.TEXT, ExtractionLimits(timeout_seconds=2.25), _fake_worker(code)
    )
    assert result.status == ae.STATUS_OK
    assert json.loads(result.segments[0].text) == [[3, 3], [0, 0]]


def test_g1_isolated_worker_ignores_a_shadow_module_beside_the_script(
    tmp_path, monkeypatch
):
    from pathlib import Path

    script = tmp_path / "attachment_extract.py"
    script.write_bytes(Path(ae.__file__).read_bytes())
    (tmp_path / "json.py").write_text("raise RuntimeError('shadow module imported')")
    monkeypatch.setattr(ae, "__file__", str(script))
    result = extract(b"isolated", ae.TEXT)
    assert result.status == ae.STATUS_OK
    assert _texts(result) == ["isolated"]


@pytest.mark.skipif(sys.platform == "win32", reason="the selector exchange is POSIX")
@pytest.mark.parametrize(
    "failure", ["timeout", "overflow", "selector_create", "selector_select"]
)
def test_g1_worker_failures_close_pipes_unregister_and_reap(monkeypatch, failure):
    import os
    import selectors
    import subprocess

    real_popen = subprocess.Popen
    children = []
    registered, unregistered = [], []

    def popen(*args, **kwargs):
        assert kwargs["stderr"] == subprocess.DEVNULL
        proc = real_popen(*args, **kwargs)
        children.append(proc)
        return proc

    class Selector(selectors.DefaultSelector):
        def __init__(self):
            if failure == "selector_create":
                raise OSError("no descriptors")
            super().__init__()

        def register(self, fd, *args):
            registered.append(fd)
            return super().register(fd, *args)

        def unregister(self, fd):
            unregistered.append(fd)
            return super().unregister(fd)

        def select(self, timeout=None):
            if failure == "selector_select":
                raise OSError("selector failed")
            return super().select(timeout)

    monkeypatch.setattr(subprocess, "Popen", popen)
    monkeypatch.setattr(selectors, "DefaultSelector", Selector)
    code = "import time; time.sleep(60)"
    if failure == "overflow":
        code = "import sys; sys.stdout.write('x' * 100000)"
    result = ae._run_worker(
        b"x",
        ae.TEXT,
        ExtractionLimits(max_chars=10, max_segments=1, timeout_seconds=0.5),
        _fake_worker(code),
    )
    expected = {
        "timeout": (ae.STATUS_TIMEOUT, "wall_clock"),
        "overflow": (ae.STATUS_FAILED, "worker_output"),
        "selector_create": (ae.STATUS_FAILED, "worker_io"),
        "selector_select": (ae.STATUS_FAILED, "worker_io"),
    }
    assert (result.status, result.detail) == expected[failure]
    assert len(children) == 1
    proc = children[0]
    assert proc.stdin.closed and proc.stdout.closed
    assert proc.returncode is not None
    assert sorted(registered) == sorted(unregistered)
    if failure != "selector_create":
        assert len(registered) == 2
    with pytest.raises(ChildProcessError):
        os.waitpid(proc.pid, os.WNOHANG)


@pytest.mark.parametrize("ctype", [ae.DOCX, ae.XLSX, ae.ODS])
def test_g7_table_rows_drop_trailing_empty_cells_but_keep_internal_cells(ctype):
    row = ["left", "", "right", "", ""]
    if ctype == ae.DOCX:
        data = _docx(paragraphs=[], table=[row])
    elif ctype == ae.XLSX:
        data = _xlsx({"Sheet": [row]})
    else:
        cells = "".join(
            f"<table:table-cell><text:p>{v}</text:p></table:table-cell>" for v in row
        )
        data = _odf(
            ae.ODS,
            f'<office:spreadsheet><table:table table:name="Sheet"><table:table-row>{cells}</table:table-row></table:table></office:spreadsheet>',
        )
    result = extract(data, ctype)
    assert result.status == ae.STATUS_OK
    assert _texts(result) == ["left\t\tright"]


def test_g7_lone_cr_and_whitespace_only_lines_preserve_block_locators():
    result = extract(b"first\r \t\rsecond\rthird", ae.TEXT)
    assert result.segments == (
        ae.Segment("first", {"line": 1}),
        ae.Segment("second\nthird", {"line": 3}),
    )


@pytest.mark.parametrize("ctype", [ae.PPTX, ae.ODP])
def test_g7_empty_presentation_paragraphs_do_not_add_separators(ctype):
    if ctype == ae.PPTX:
        data = _pptx([["", "first", "  ", "second", ""]])
    else:
        paragraphs = "".join(
            f"<text:p>{v}</text:p>" for v in ["", "first", "  ", "second", ""]
        )
        data = _odf(
            ae.ODP,
            f"<office:presentation><draw:page>{paragraphs}</draw:page></office:presentation>",
        )
    assert _texts(extract(data, ctype)) == ["first\nsecond"]


def test_g7_odf_line_break_zero_spaces_and_nested_paragraphs():
    body = '<office:text><text:p>outer<text:line-break/>line<text:s text:c="0"/>end<draw:frame><text:p>nested</text:p></draw:frame>tail</text:p></office:text>'
    result = extract(_odf(ae.ODT, body), ae.ODT)
    assert _texts(result) == ["outer\nline endtail", "nested"]


def test_g7_ods_covered_cells_keep_column_positions():
    body = '<office:spreadsheet><table:table table:name="S"><table:table-row><table:table-cell><text:p>a</text:p></table:table-cell><table:covered-table-cell/><table:table-cell><text:p>b</text:p></table:table-cell></table:table-row></table:table></office:spreadsheet>'
    assert _texts(extract(_odf(ae.ODS, body), ae.ODS)) == ["a\t\tb"]


def test_g7_xlsx_ignores_phonetic_guides_and_negative_shared_indices():
    data = _xlsx({"": [["s:0", "s:-1", "tail"]]})
    data = _replace_member(
        data,
        "xl/sharedStrings.xml",
        f'<sst xmlns="{_S}"><si><t>word</t><rPh><t>phonetic</t></rPh></si></sst>',
    )
    result = extract(data, ae.XLSX)
    assert result.segments == (ae.Segment("word\t\ttail", {"sheet": ""}),)


@pytest.mark.parametrize("kind", ["external", "absolute", "missing"])
def test_g7_xlsx_relationships_respect_external_absolute_and_missing_targets(kind):
    data = _xlsx({"S": [["visible"]]})
    target = (
        "/xl/worksheets/sheet1.xml"
        if kind != "missing"
        else "/xl/worksheets/missing.xml"
    )
    rels = _rels([("rId1", _REL_WS, target)])
    if kind == "external":
        rels = rels.replace('Id="rId1"', 'Id="rId1" TargetMode="External"')
    result = extract(_replace_member(data, "xl/_rels/workbook.xml.rels", rels), ae.XLSX)
    assert result.status == ae.STATUS_OK
    assert _texts(result) == (["visible"] if kind == "absolute" else [])


def test_g7_docx_body_custom_xml_wraps_paragraphs():
    data = _docx(["wrapped"])
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        body = archive.read("word/document.xml").decode()
    body = body.replace("<w:body>", "<w:body><w:customXml>").replace(
        "</w:body>", "</w:customXml></w:body>"
    )
    assert _texts(
        extract(_replace_member(data, "word/document.xml", body), ae.DOCX)
    ) == ["wrapped"]


def test_g7_docx_cell_paragraphs_are_joined_by_newlines():
    document = docx.Document(io.BytesIO(_docx(paragraphs=[], table=[["first"]])))
    document.tables[0].cell(0, 0).add_paragraph("second")
    out = io.BytesIO()
    document.save(out)
    assert _texts(extract(out.getvalue(), ae.DOCX)) == ["first\nsecond"]


def test_g7_docx_text_box_at_paragraph_start_has_no_leading_newline():
    data = _docx(["after"])
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        body = archive.read("word/document.xml").decode()
    body = body.replace(
        "<w:r><w:t>after</w:t></w:r>",
        "<w:r><w:pict><w:txbxContent><w:p><w:r><w:t>box</w:t></w:r></w:p></w:txbxContent></w:pict></w:r><w:r><w:t>after</w:t></w:r>",
    )
    assert _texts(
        extract(_replace_member(data, "word/document.xml", body), ae.DOCX)
    ) == ["box\nafter"]


def test_g3_docx_mixed_case_image_content_type_is_opaque():
    ae._refuse_docx_dtds(
        {
            "[Content_Types].xml": b'<Types><Override PartName="/word/picture.xml" ContentType="ImAgE/svg+xml"/></Types>',
            "word/picture.xml": b"<!DOCTYPE svg><svg/>",
        }
    )


def test_g4_segment_after_exactly_full_budget_adds_no_empty_segment():
    out = ae._Collector(4, max_segments=10)
    assert out.add("full", {"page": 1})
    assert not out.add("extra", {"page": 2})
    assert out.segments == [ae.Segment("full", {"page": 1})]
    assert out.truncated


def test_g4_pdf_text_with_a_lone_surrogate_is_ok(monkeypatch):
    import pypdf
    from types import SimpleNamespace

    monkeypatch.setattr(
        pypdf,
        "PdfReader",
        lambda _: SimpleNamespace(
            is_encrypted=False, pages=[SimpleNamespace(extract_text=lambda: "a\ud800b")]
        ),
    )
    result = ae._extract_in_process(b"pdf", ae.PDF, ExtractionLimits())
    assert result.status == ae.STATUS_OK
    assert _texts(result) == ["a?b"]
    result.segments[0].text.encode("utf-8")
